from __future__ import annotations

from pathlib import Path

import openpyxl
import pypdfium2 as pdfium
import xlrd
from pypdf import PdfReader


class DocumentTextExtractor:
    SUPPORTED_EXTENSIONS = {".pdf", ".xls", ".xlsx"}

    def extract(self, file_path: str) -> str:
        path = Path(file_path)
        extension = path.suffix.lower()
        if extension == ".pdf":
            return self._extract_pdf(path)
        if extension == ".xlsx":
            return self._extract_xlsx(path)
        if extension == ".xls":
            return self._extract_xls(path)
        raise ValueError("unsupported_text_document_type")

    def _extract_pdf(self, path: Path) -> str:
        reader = PdfReader(str(path))
        chunks: list[str] = []
        for index, page in enumerate(reader.pages, start=1):
            # PDF drawing order can put a customer's name beside a supplier's
            # requisites. Preserve visual rows before interpreting party roles.
            text = (page.extract_text(extraction_mode="layout") or "").strip()
            if text:
                chunks.append(f"[Page {index}]\n{text}")
        return "\n\n".join(chunks)

    def render_pdf_pages(self, file_path: str, output_dir: str, max_pages: int = 10, scale: float = 3.0) -> list[Path]:
        path = Path(file_path)
        document = pdfium.PdfDocument(str(path))
        image_paths: list[Path] = []
        if len(document) > max_pages:
            raise ValueError("source_document_too_many_pages")
        page_count = len(document)
        for index in range(page_count):
            page = document[index]
            bitmap = page.render(scale=scale)
            image = bitmap.to_pil()
            output_path = Path(output_dir) / f"{path.stem}-page-{index + 1}.png"
            image.save(output_path, format="PNG")
            image_paths.append(output_path)
        return image_paths

    def _extract_xlsx(self, path: Path) -> str:
        workbook = openpyxl.load_workbook(path, data_only=True)
        chunks: list[str] = []
        for sheet in workbook.worksheets:
            rows: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                values = [str(value).strip() for value in row if value not in (None, "")]
                if values:
                    rows.append(" | ".join(values))
            if rows:
                chunks.append(f"[Sheet: {sheet.title}]\n" + "\n".join(rows))
        return "\n\n".join(chunks)

    def _extract_xls(self, path: Path) -> str:
        workbook = xlrd.open_workbook(path)
        chunks: list[str] = []
        for sheet in workbook.sheets():
            rows: list[str] = []
            for row_index in range(sheet.nrows):
                values = [str(value).strip() for value in sheet.row_values(row_index) if str(value).strip()]
                if values:
                    rows.append(" | ".join(values))
            if rows:
                chunks.append(f"[Sheet: {sheet.name}]\n" + "\n".join(rows))
        return "\n\n".join(chunks)
